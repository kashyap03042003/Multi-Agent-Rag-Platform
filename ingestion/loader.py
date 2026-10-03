from pathlib import Path
from bs4 import BeautifulSoup
from docx import Document
from pptx import Presentation
from pypdf import PdfReader

SKIP_DIRS = {"processed"}


def read_pdf(path: Path) -> list[tuple]:
    return [(i, p.extract_text() or "") for i, p in enumerate(PdfReader(path).pages, start=1)]


def read_docx(path: Path) -> list[tuple]:
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return [(None, "\n".join(parts))]


def read_pptx(path: Path) -> list[tuple]:
    slides = []
    for i, slide in enumerate(Presentation(path).slides, start=1):
        texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
        slides.append((i, "\n".join(texts)))
    return slides


def read_html(path: Path) -> list[tuple]:
    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
    for tag in soup(["script", "style", "nav", "footer"]):
        tag.decompose()
    return [(None, soup.get_text(separator="\n"))]


def read_txt(path: Path) -> list[tuple]:
    return [(None, path.read_text(encoding="utf-8", errors="ignore"))]


READERS = {
    ".pdf": read_pdf, ".docx": read_docx, ".pptx": read_pptx,
    ".html": read_html, ".htm": read_html, ".txt": read_txt,
}


def is_supported(rel: Path) -> bool:
    is_temp = rel.name.startswith((".", "~$"))
    return rel.suffix.lower() in READERS and not is_temp and not (SKIP_DIRS & set(rel.parts))


def load_file(path: Path, root: Path) -> list[dict]:
    rel = path.relative_to(root)
    sections = READERS[path.suffix.lower()](path)
    return [
        {
            "text": text,
            "source": str(rel),
            "category": rel.parts[0],
            "format": path.suffix.lower().lstrip("."),
            "page": page,
        }
        for page, text in sections
        if text.strip()
    ]


def load_documents(data_dir: str) -> list[dict]:
    root = Path(data_dir)
    pages = []
    for path in sorted(root.rglob("*")):
        rel = path.relative_to(root)
        if not is_supported(rel):
            continue
        try:
            file_pages = load_file(path, root)
        except Exception as e:
            print(f"Skipped {rel}: {e}")
            continue
        pages.extend(file_pages)
        print(f"Loaded {rel} ({len(file_pages)} sections)")
    return pages
